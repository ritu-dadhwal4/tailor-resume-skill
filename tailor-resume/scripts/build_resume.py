#!/usr/bin/env python3
"""Build an ATS-friendly, colour-highlighted resume (.docx) from a JSON spec.

Usage:
  python3 build_resume.py spec.json out.docx [--highlights y,g,b] [--pages 1]
                          [--paper a4|letter] [--legend]

Inline highlight markup inside any text field:
  {y:...}  yellow - behaviour you proved
  {g:...}  green  - achievements in numbers
  {b:...}  blue   - the job's own language
Codes left out of --highlights are written as plain text.
Placeholders such as [X%] are allowed but reported, so they get filled before sending.
"""
import argparse
import json
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor, Mm
except ImportError:
    sys.exit("python-docx is missing. Run: python3 -m pip install --user python-docx")

COLOURS = {"y": "FFF1A6", "g": "CDEFD0", "b": "D3E5FF"}
LABELS = {"y": "Behaviour you proved", "g": "Achievements in numbers", "b": "The job's language"}
FONT = "Calibri"
TEXT = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x55, 0x55, 0x55)
MARK = re.compile(r"\{([ygb]):(.+?)\}")
PLACEHOLDER = re.compile(r"\[[^\]]*\b(X|N|\?)\b[^\]]*\]")

# Rough capacity for the page-length warning (lines of 10.5pt text).
CHARS_PER_LINE = {"a4": 112, "letter": 116}
LINES_PER_PAGE = 54


def segments(text, enabled):
    """Split text into (chunk, colour_code_or_None)."""
    out, pos = [], 0
    for m in MARK.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], None))
        code = m.group(1)
        out.append((m.group(2), code if code in enabled else None))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], None))
    return out


def plain(text):
    return MARK.sub(lambda m: m.group(2), text or "")


def shade(run, hex_fill):
    rpr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    rpr.append(shd)


def style_run(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_marked(par, text, enabled, size=10.5, bold=False, color=TEXT):
    for chunk, code in segments(text, enabled):
        r = par.add_run(chunk)
        style_run(r, size=size, bold=bold, color=color)
        if code:
            shade(r, COLOURS[code])


def spacing(par, before=0, after=0, line=1.08):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def bottom_border(par):
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "444444")
    bdr.append(b)
    ppr.append(bdr)


class Builder:
    def __init__(self, enabled, paper):
        self.doc = Document()
        self.enabled = enabled
        self.lines = 0
        self.cpl = CHARS_PER_LINE[paper]
        sec = self.doc.sections[0]
        if paper == "a4":
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.7))
        sec.top_margin = sec.bottom_margin = Inches(0.6)
        self.right_tab = sec.page_width - sec.left_margin - sec.right_margin
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10.5)

    def count(self, text, extra=0.0):
        self.lines += max(1, -(-len(plain(text)) // self.cpl)) + extra

    def header(self, name, contact):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(p.add_run(name), size=18, bold=True)
        spacing(p, after=2)
        if contact:
            line = "  |  ".join(contact)
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(p.add_run(line), size=9.5, color=MUTED)
            spacing(p, after=6)
        self.lines += 3

    def heading(self, title):
        p = self.doc.add_paragraph()
        style_run(p.add_run(title.upper()), size=11, bold=True)
        spacing(p, before=10, after=4)
        bottom_border(p)
        self.lines += 2

    def para(self, text, align_center=False):
        p = self.doc.add_paragraph()
        if align_center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_marked(p, text, self.enabled)
        spacing(p, after=2)
        self.count(text)

    def bullet(self, text, label=None):
        p = self.doc.add_paragraph(style="List Bullet")
        if label:
            r = p.add_run(label + ": ")
            style_run(r, bold=True)
        add_marked(p, text, self.enabled)
        spacing(p, after=2)
        p.paragraph_format.left_indent = Inches(0.22)
        self.count((label + ": " if label else "") + text, 0.1)

    def row(self, left_bold, left_rest, right, italic_line=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(self.right_tab, WD_TAB_ALIGNMENT.RIGHT)
        add_marked(p, left_bold, self.enabled, bold=True)
        if left_rest:
            add_marked(p, left_rest, self.enabled)
        if right:
            r = p.add_run("\t" + right)
            style_run(r, color=MUTED)
        spacing(p, before=6, after=1)
        self.lines += 1.3
        if italic_line:
            p = self.doc.add_paragraph()
            add_marked(p, italic_line, self.enabled)
            spacing(p, after=1)
            self.lines += 1

    def legend(self):
        codes = [c for c in "ygb" if c in self.enabled]
        if not codes:
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, before=10)
        for i, c in enumerate(codes):
            if i:
                style_run(p.add_run("     "), size=9)
            r = p.add_run(" " + LABELS[c] + " ")
            style_run(r, size=9, color=MUTED)
            shade(r, COLOURS[c])
        self.lines += 2


def build(spec, enabled, paper, legend):
    b = Builder(enabled, paper)
    b.header(spec["name"], spec.get("contact", []))

    def summary():
        if spec.get("summary"):
            b.heading(spec.get("summary_heading", "Summary"))
            b.para(spec["summary"])

    def skills():
        if spec.get("skills"):
            b.heading("Skills")
            for s in spec["skills"]:
                b.bullet(s["items"], s.get("label"))

    def experience():
        if spec.get("experience"):
            b.heading("Experience")
            for job in spec["experience"]:
                rest = " - " + job["title"] if job.get("title") else ""
                if job.get("location"):
                    rest += ", " + job["location"]
                b.row(job["company"], rest, job.get("dates", ""))
                for x in job.get("bullets", []):
                    b.bullet(x)

    def projects():
        if spec.get("projects"):
            b.heading("Projects")
            for pr in spec["projects"]:
                b.bullet(pr["text"], pr.get("name"))

    def education():
        if spec.get("education"):
            b.heading("Education")
            for e in spec["education"]:
                b.row(e["school"], (", " + e["location"]) if e.get("location") else "",
                      e.get("dates", ""), e.get("degree"))
                for x in e.get("bullets", []):
                    b.bullet(x)

    def certifications():
        if spec.get("certifications"):
            b.heading("Certifications")
            for c in spec["certifications"]:
                b.bullet(c)

    def extras():
        for sec in spec.get("extra_sections", []):
            b.heading(sec["heading"])
            for x in sec.get("bullets", []):
                b.bullet(x)

    parts = {"summary": summary, "skills": skills, "experience": experience,
             "projects": projects, "education": education,
             "certifications": certifications, "extra": extras}
    order = spec.get("section_order",
                     ["summary", "skills", "experience", "projects",
                      "education", "certifications", "extra"])
    for key in order:
        parts[key]()
    if legend:
        b.legend()
    return b


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("out")
    ap.add_argument("--highlights", default="y,g,b",
                    help="comma list of y,g,b or 'none'")
    ap.add_argument("--pages", type=int, default=1)
    ap.add_argument("--paper", choices=["a4", "letter"], default="a4")
    ap.add_argument("--legend", action="store_true")
    a = ap.parse_args()

    enabled = set() if a.highlights == "none" else set(a.highlights.split(","))
    spec = json.load(open(a.spec, encoding="utf-8"))
    b = build(spec, enabled, a.paper, a.legend)
    b.doc.save(a.out)

    raw = json.dumps(spec, ensure_ascii=False)
    holes = PLACEHOLDER.findall(raw)
    est = b.lines / LINES_PER_PAGE
    print(f"saved: {a.out}")
    print(f"estimated length: {est:.2f} pages (target {a.pages})")
    if est > a.pages + 0.02:
        print(f"WARNING: likely over {a.pages} page(s) - cut ~{int((est - a.pages) * LINES_PER_PAGE) + 1} lines")
    elif a.pages == 2 and est < 1.3:
        print("NOTE: content fits on one page - consider --pages 1")
    counts = {c: len(re.findall(r"\{" + c + r":", raw)) for c in "ygb"}
    print("highlights: " + ", ".join(f"{LABELS[c]}={counts[c]}{'' if c in enabled else ' (off)'}" for c in "ygb"))
    if holes:
        print(f"PLACEHOLDERS TO FILL: {len(PLACEHOLDER.findall(raw))} (search the doc for '[')")


if __name__ == "__main__":
    main()
