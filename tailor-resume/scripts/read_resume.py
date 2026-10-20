#!/usr/bin/env python3
"""Print the text of a .docx resume (PDFs and images are read directly by Claude)."""
import sys
from docx import Document

d = Document(sys.argv[1])
for p in d.paragraphs:
    if p.text.strip():
        print(p.text)
for t in d.tables:
    for row in t.rows:
        print(" | ".join(c.text.strip() for c in row.cells))
